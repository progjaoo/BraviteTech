#!/usr/bin/env python3
"""Recria a apresentação em PPTX/PDF usando textos e formas editáveis.
Dependências: python-pptx, reportlab, pymupdf, Pillow, inkscape.
"""
from pathlib import Path
import os, re, subprocess, math
import zipfile
import xml.etree.ElementTree as ET
from io import BytesIO
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
import fitz
from PIL import Image, ImageDraw

BASE = Path(__file__).resolve().parent
FONTS = BASE.parent / '05-fontes' / 'desktop'
RES = BASE / 'recursos'
RES.mkdir(exist_ok=True)
W,H = 960,540
C = {'black':'0A0A0A','blue':'2D6BFF','white':'FFFFFF','muted':'5A6475','line':'D9DEE7','darkmuted':'B6BDC9'}
for weight in (400,500,700):
    pdfmetrics.registerFont(TTFont('Space'+str(weight), str(FONTS / f'SpaceGrotesk-{weight}.ttf')))

src = (RES/'logo-original-opcao03.svg').read_text()
for variant,fg in [('principal',C['black']),('clara',C['white'])]:
    svg=re.sub(r'width="780" height="140" viewBox="45 62 780 140"', 'width="760" height="110" viewBox="58 76 760 110"', src)
    svg=svg.replace('#0A0A0A','#'+fg)
    (RES/f'logo-{variant}.svg').write_text(svg)
    for ext in ('png','pdf'):
        cmd=['inkscape',str(RES/f'logo-{variant}.svg'),f'--export-filename={RES/f"logo-{variant}.{ext}"}']
        if ext=='png': cmd += ['--export-width=2496']
        subprocess.run(cmd,check=True,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)

ppt=Presentation(); ppt.slide_width=Inches(W/72); ppt.slide_height=Inches(H/72)
pdf=canvas.Canvas(str(RES/'slides-base.pdf'),pagesize=(W,H))
pdf.setTitle('Bravite — Apresentação institucional e comercial | Modelo editável')
pdf.setAuthor('Bravite')
logo_placements=[]
slide=None
page=0

def color(c): return C.get(c,c)
def rgb(c): return RGBColor.from_string(color(c))
def begin(bg='white'):
    global slide,page
    page+=1
    slide=ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=rgb(bg)
    pdf.setFillColor(HexColor('#'+color(bg))); pdf.rect(0,0,W,H,fill=1,stroke=0)

def rect(x,y,w,h,fill,r=0):
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x/72), Inches(y/72), Inches(w/72), Inches(h/72))
    shape.fill.solid(); shape.fill.fore_color.rgb=rgb(fill); shape.line.fill.background()
    pdf.setFillColor(HexColor('#'+color(fill))); pdf.rect(x,H-y-h,w,h,fill=1,stroke=0)

def line(x1,y1,x2,y2,c='line',width=1):
    from pptx.enum.shapes import MSO_CONNECTOR
    shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1/72),Inches(y1/72),Inches(x2/72),Inches(y2/72))
    shape.line.color.rgb=rgb(c); shape.line.width=Pt(width)
    pdf.setStrokeColor(HexColor('#'+color(c))); pdf.setLineWidth(width); pdf.line(x1,H-y1,x2,H-y2)

def txt(x,y,text,size=20,weight=400,c='black',width=840,lh=None):
    lh=lh or size*1.22
    rows=text.split('\n')
    for row in rows:
        measured=pdfmetrics.stringWidth(row,'Space'+str(weight),size)
        if measured>width+3: print('TEXT OVERFLOW',page,repr(row),round(measured),width)
    box=slide.shapes.add_textbox(Inches(x/72),Inches(y/72),Inches(width/72),Inches((len(rows)*lh+10)/72))
    tf=box.text_frame; tf.clear(); tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,row in enumerate(rows):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=row; p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=Pt(lh)
        p.font.name='Space Grotesk'; p.font.size=Pt(size); p.font.bold=weight==700; p.font.color.rgb=rgb(c)
        if weight==500: p.font.name='Space Grotesk Medium'
    pdf.setFont('Space'+str(weight),size); pdf.setFillColor(HexColor('#'+color(c)))
    # True font ascender aligns ReportLab to PowerPoint text box's top baseline.
    asc=pdfmetrics.getAscent('Space'+str(weight),size)
    for i,row in enumerate(rows): pdf.drawString(x,H-y-asc-i*lh,row)
    return box

def logo(x,y,w=200,variant='principal'):
    h=w*110/760
    slide.shapes.add_picture(str(RES/f'logo-{variant}.png'),Inches(x/72),Inches(y/72),width=Inches(w/72),height=Inches(h/72))
    logo_placements.append((page-1,x,y,w,h,variant))

def footer(num,dark=False,note=''):
    ink='darkmuted' if dark else 'muted'
    line(60,486,900,486,'454B55' if dark else 'line',.6)
    txt(60,503,'BRAVITE  /  SOFTWARE HOUSE',10,500,ink)
    if note: txt(345,503,note,10,400,ink,width=475)
    txt(868,499,f'{num:02d}',16,500,'white' if dark else 'black',width=32)

def notes(text):
    slide.notes_slide.notes_text_frame.text=text

def end(): pdf.showPage()

# 01 / Cover
begin('black')
logo(60,49,228,'clara')
txt(650,54,'APRESENTAÇÃO INSTITUCIONAL',11,500,'darkmuted',width=255)
txt(60,169,'Coragem para criar.\nEngenharia para evoluir.',48,400,'white',width=810,lh=62)
rect(60,370,166,34,'blue')
txt(74,379,'SOFTWARE HOUSE',12,700,'white',width=145)
txt(60,433,'www.bravite.com.br',15,400,'white',width=350)
txt(697,433,'Instagram  @bravite.br',15,400,'white',width=240)
line(843,169,902,169,'blue',6); line(843,199,885,199,'blue',6); line(843,229,902,229,'blue',6)
notes('Apresentação institucional e modelo comercial. Assinatura de marca sugerida. O nome e os contatos são os dados informados pelo proprietário. Os slides 3, 4 e 5 contêm sugestões ou campos a validar antes de envio comercial.')
end()

# 02 / Meaning
begin()
txt(60,47,'01 / ESSÊNCIA',11,500,'muted')
txt(60,110,'Um nome.\nTrês ideias.',48,400,width=365,lh=57)
txt(60,283,'A Bravite é uma software house\ncujo nome reúne coragem,\ntecnologia e engenharia.',18,400,'muted',width=350,lh=27)
line(439,112,439,420)
rows=[(115,'BRAVE','Coragem.'),(222,'IT','Information Technology.'),(329,'E','Engineer.')]
for y,title,desc in rows:
    txt(481,y,title,35,500,width=185)
    txt(674,y+12,desc,17,400,'muted',width=240)
    line(480,y+78,900,y+78)
footer(2)
notes('Origem do nome fornecida pelo proprietário: BRAVE = coragem; IT = Information Technology; E = Engineer. Não apresentar essas partes como sigla técnica ou certificação.')
end()

# 03 / proposed catalog
begin('black')
txt(60,47,'02 / FRENTES DE ATUAÇÃO PROPOSTAS',11,500,'darkmuted')
txt(60,101,'Tecnologia com\nintenção de construir.',43,400,'white',width=780,lh=53)
cards=[('01','Experiências web','Interfaces e produtos\ndigitais para a web.'),('02','Sistemas sob medida','Ferramentas para necessidades\nespecíficas de uma operação.'),('03','Integrações','Conexões entre sistemas\ne fluxos de informação.')]
for i,(n,title,desc) in enumerate(cards):
    x=60+i*286
    rect(x,266,268,182,'white')
    txt(x+20,286,n,12,700,'black',width=200)
    txt(x+20,328,title,20,500,'black',width=235)
    txt(x+20,374,desc,15,400,'muted',width=235,lh=23)
footer(3,True,'CATÁLOGO SUGERIDO · VALIDAR ANTES DE PUBLICAR')
notes('MODELO. As três frentes foram propostas para organizar uma possível oferta. Elas não constituem catálogo oficial confirmado. Validar capacidade, escopo, tecnologias, manutenção e suporte. Remover esta nota e a indicação de modelo apenas depois da validação.')
end()

# 04 / Workflow
begin()
txt(60,47,'03 / COLABORAÇÃO',11,500,'muted')
txt(60,108,'Da ideia à evolução.',46,400,width=840)
txt(60,176,'Um caminho de colaboração proposto.',18,400,'muted')
steps=[('01','Entender','Contexto, pessoas,\nnecessidades e restrições.'),('02','Estruturar','Escopo, prioridades\ne critérios de aceite.'),('03','Construir','Implementação e\ndecisões acompanhadas.'),('04','Evoluir','Revisão da entrega\ne próximos passos.')]
line(81,277,737,277,'line',2)
for i,(n,title,desc) in enumerate(steps):
    x=60+i*218
    rect(x,256,44,44,'blue')
    txt(x+10,268,n,15,700,'white',width=28)
    txt(x,327,title,26,500,width=203)
    txt(x,377,desc,15,400,'muted',width=202,lh=24)
footer(4,False,'PROCESSO SUGERIDO · ADAPTAR À OPERAÇÃO')
notes('MODELO. Processo sugerido, a validar com a operação. Não representa SLA, prazo, número de reuniões ou obrigação comercial. Ajustar à realidade e ao contrato de cada projeto.')
end()

# 05 / Template
begin()
txt(60,47,'04 / BASE DA PROPOSTA',11,500,'muted')
txt(685,47,'MODELO PARA PREENCHER',11,500,'muted',width=225)
txt(60,107,'[Cliente / projeto]',44,400,width=820)
fields=[('OBJETIVO','[Resultado desejado]'),('ENTREGÁVEIS E LIMITES','[Escopo incluído / excluído]'),('ETAPAS E PRAZOS','[Cronograma acordado]'),('INVESTIMENTO E CONDIÇÕES','[Valor / condições de pagamento]'),('CRITÉRIOS DE ACEITE','[Como validar a entrega]'),('RESPONSABILIDADES','[Acessos / insumos / aprovações]')]
for i,(lab,value) in enumerate(fields):
    x=60+(i%2)*432; y=202+(i//2)*88
    line(x,y,x+408,y,'line',1)
    txt(x,y+14,lab,10,700,'muted',width=400)
    txt(x,y+39,value,18,400,'black',width=408)
footer(5,False,'SUBSTITUIR OS CAMPOS ANTES DO ENVIO')
notes('Preencher todos os campos entre colchetes com dados do projeto. Confirmar escopo, exclusões, premissas, acessos, etapas, prazos, investimento, tributos aplicáveis, pagamento, responsabilidades e aceite. O slide é um resumo; não substitui proposta detalhada ou contrato. Não enviar com campos em branco.')
end()

# 06 / Contact
begin('black')
logo(60,49,228,'clara')
txt(60,169,'Vamos dar forma\nà próxima ideia.',52,400,'white',width=835,lh=64)
rect(60,363,47,5,'blue')
txt(60,399,'www.bravite.com.br',22,500,'white',width=490)
txt(60,440,'Instagram  @bravite.br',17,400,'darkmuted',width=490)
txt(673,415,'CORAGEM. TECNOLOGIA.\nENGENHARIA.',12,500,'white',width=240,lh=19)
notes('Contatos fornecidos pelo proprietário: www.bravite.com.br e Instagram @bravite.br. Conferir a disponibilidade dos canais antes de distribuir. Não foi informado e-mail ou telefone.')
end()

pdf.save()
ppt.core_properties.title='Bravite — Apresentação institucional e comercial'
ppt.core_properties.subject='Modelo editável: essência, oferta sugerida, processo e proposta'
ppt.core_properties.author='Bravite'
ppt.core_properties.keywords='Bravite, software house, apresentação, modelo'
ppt.save(BASE/'bravite-apresentacao-editavel.pptx')

# A paleta e a tipografia do tema também acompanham o rebrand em novas formas/slides.
deck=BASE/'bravite-apresentacao-editavel.pptx'
updated=BASE/'bravite-apresentacao-tema.pptx'
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
theme_colors={'dk1':'0A0A0A','lt1':'FFFFFF','dk2':'5A6475','lt2':'F3F4F6','accent1':'2D6BFF','accent2':'0A0A0A','accent3':'5A6475','accent4':'D9DEE7','accent5':'F3F4F6','accent6':'FFFFFF','hlink':'2D6BFF','folHlink':'0A0A0A'}
with zipfile.ZipFile(deck) as srczip, zipfile.ZipFile(updated,'w',zipfile.ZIP_DEFLATED) as destzip:
    for item in srczip.infolist():
        data=srczip.read(item.filename)
        if item.filename.startswith('ppt/theme/') and item.filename.endswith('.xml'):
            root=ET.fromstring(data)
            scheme=root.find('.//a:clrScheme',ns)
            if scheme is not None:
                scheme.set('name','Bravite 2.0')
                for element in scheme:
                    role=element.tag.split('}')[-1]
                    if role in theme_colors:
                        for child in list(element):element.remove(child)
                        ET.SubElement(element,'{'+ns['a']+'}srgbClr',{'val':theme_colors[role]})
            for latin in root.findall('.//a:fontScheme//a:latin',ns):latin.set('typeface','Space Grotesk')
            data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
        destzip.writestr(item,data)
updated.replace(deck)

# Merge true vector artwork into PDF, preserving original paths.
doc=fitz.open(RES/'slides-base.pdf')
for n,x,y,w,h,var in logo_placements:
    mark=fitz.open(RES/f'logo-{var}.pdf')
    doc[n].show_pdf_page(fitz.Rect(x,y,x+w,y+h),mark,0)
    mark.close()
doc.set_metadata({'title':'Bravite — Apresentação institucional e comercial','author':'Bravite','subject':'Modelo editável: validar oferta, processo e campos comerciais'})
doc.save(BASE/'bravite-apresentacao.pdf',garbage=4,deflate=True)
# Multi-slide preview for navigation/review, not a production logo.
thumbs=[]
for p in doc:
    pix=p.get_pixmap(matrix=fitz.Matrix(1,1),alpha=False)
    thumbs.append(Image.open(BytesIO(pix.tobytes('png'))).convert('RGB'))
board=Image.new('RGB',(1992,1734),'#F3F4F6')
for i,im in enumerate(thumbs):
    x=24+(i%2)*984; y=24+(i//2)*570
    board.paste(im,(x,y))
board.save(BASE/'previa-apresentacao.jpg',quality=93)
doc.close()
(RES/'slides-base.pdf').unlink(missing_ok=True)
print('Criados: PPTX editável, PDF vetorial, prévia de 6 slides.')
